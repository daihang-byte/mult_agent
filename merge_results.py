import os
import pandas as pd
import json
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
import logging

# 配置日志
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)


def merge_experiment_results(results_dir="results", output_file="merged_results.xlsx"):
    """
    合并所有实验文件夹的结果到单个Excel文件
    :param results_dir: 结果目录路径
    :param output_file: 输出Excel文件名
    """
    # 创建Excel工作簿
    wb = Workbook()

    # 初始化数据收集器
    all_model_data = []
    all_agent_data = []
    all_events = []
    all_decisions = []

    # 遍历结果目录
    experiment_folders = [f for f in os.listdir(results_dir)
                          if os.path.isdir(os.path.join(results_dir, f)) and f.startswith("experiment_")]

    if not experiment_folders:
        logger.error(f"在 {results_dir} 中没有找到实验文件夹")
        return

    logger.info(f"找到 {len(experiment_folders)} 个实验文件夹，开始合并数据...")

    for exp_folder in experiment_folders:
        exp_path = os.path.join(results_dir, exp_folder)

        try:
            # 读取元数据
            meta_path = os.path.join(exp_path, "meta.json")
            with open(meta_path, 'r', encoding='utf-8') as f:
                meta = json.load(f)

            transparency = meta.get('transparent', 'unknown')
            seed = meta.get('seed', 'unknown')

            # 添加实验标识列
            exp_id = exp_folder.replace("experiment_", "")

            # 处理模型数据
            model_path = os.path.join(exp_path, "model_data.csv")
            if os.path.exists(model_path):
                df_model = pd.read_csv(model_path)
                df_model['experiment'] = exp_id
                df_model['transparency'] = transparency
                df_model['seed'] = seed
                all_model_data.append(df_model)

            # 处理代理数据
            agent_path = os.path.join(exp_path, "agent_data.csv")
            if os.path.exists(agent_path):
                df_agent = pd.read_csv(agent_path)
                df_agent['experiment'] = exp_id
                df_agent['transparency'] = transparency
                df_agent['seed'] = seed
                all_agent_data.append(df_agent)

            # 处理事件数据
            events_path = os.path.join(exp_path, "events.csv")
            if os.path.exists(events_path):
                df_events = pd.read_csv(events_path)
                df_events['experiment'] = exp_id
                all_events.append(df_events)

            # 处理决策数据
            decisions_path = os.path.join(exp_path, "decisions.csv")
            if os.path.exists(decisions_path):
                df_decisions = pd.read_csv(decisions_path)
                df_decisions['experiment'] = exp_id
                all_decisions.append(df_decisions)

            logger.info(f"已处理: {exp_folder}")

        except Exception as e:
            logger.error(f"处理 {exp_folder} 时出错: {str(e)}")

    # 合并所有数据
    logger.info("开始合并数据...")
    merged_model = pd.concat(all_model_data, ignore_index=True) if all_model_data else None
    merged_agent = pd.concat(all_agent_data, ignore_index=True) if all_agent_data else None
    merged_events = pd.concat(all_events, ignore_index=True) if all_events else None
    merged_decisions = pd.concat(all_decisions, ignore_index=True) if all_decisions else None

    # 保存到Excel的不同sheet
    logger.info(f"保存结果到 {output_file}")

    # 删除默认创建的空sheet
    if 'Sheet' in wb.sheetnames:
        del wb['Sheet']

    # 添加各数据表到不同sheet
    if merged_model is not None:
        ws_model = wb.create_sheet("Model Data")
        for r in dataframe_to_rows(merged_model, index=False, header=True):
            ws_model.append(r)

    if merged_agent is not None:
        ws_agent = wb.create_sheet("Agent Data")
        for r in dataframe_to_rows(merged_agent, index=False, header=True):
            ws_agent.append(r)

    if merged_events is not None:
        ws_events = wb.create_sheet("Events")
        for r in dataframe_to_rows(merged_events, index=False, header=True):
            ws_events.append(r)

    if merged_decisions is not None:
        ws_decisions = wb.create_sheet("Decisions")
        for r in dataframe_to_rows(merged_decisions, index=False, header=True):
            ws_decisions.append(r)

    # 保存Excel文件
    wb.save(output_file)
    logger.info(f"数据合并完成! 结果保存到: {output_file}")


# 执行合并
if __name__ == "__main__":
    # 配置参数
    RESULTS_DIR = "results"  # 修改为你的结果目录
    OUTPUT_FILE = "merged_experiment_results1.xlsx"

    merge_experiment_results(RESULTS_DIR, OUTPUT_FILE)




import pandas as pd
# 1. 加载 Excel 文件中的 Agent Data 表
file_path = r"D:\乐云云盘\thinking\merged_experiment_results.xlsx"
sheet_name = "Agent Data"

df = pd.read_excel(file_path, sheet_name=sheet_name)

# 2. 创建 moral_group 列：高道德为1，低道德为0，分界点为0.6
df['moral_group'] = (df['moral_value'] >= 0.6).astype(int)

# 3. 可选：保存为新的 Excel 或 CSV 文件
output_path = r"D:\乐云云盘\thinking\agent_data_with_moral_group.xlsx"
df.to_excel(output_path, index=False)

# 如果你更倾向于 CSV，可改为：
# df.to_csv(r"D:\乐云云盘\thinking\agent_data_with_moral_group.csv", index=False)

print("✅ moral_group 划分完成，文件已保存：", output_path)





import pandas as pd
import numpy as np

def calculate_gini(array):
    """计算基尼系数"""
    array = np.array(array)
    if array.size == 0:
        return 0
    array = np.sort(array)
    n = array.size
    cum_wealth = np.cumsum(array)
    gini = (2 * np.sum((np.arange(1, n+1)) * array)) / (n * np.sum(array)) - (n + 1) / n
    return gini

# 读取 Excel 文件
file_path = r"D:\word\工作簿3低道德.xlsx"
df = pd.read_excel(file_path, sheet_name='Sheet1')

# 按 experiment 和 step 分组，计算基尼系数和总财富
result = (
    df.groupby(['experiment', 'step'])
    .agg(
        gini=('wealth', calculate_gini),
        total_wealth=('wealth', 'sum')
    )
    .reset_index()
)

# 保存结果为新 Excel 文件
output_path = r"D:\word\低道德_gini_total_wealth.xlsx"
result.to_excel(output_path, index=False)

print(f"处理完成，结果已保存为：{output_path}")


import pandas as pd
import numpy as np

def calculate_gini(array):
    """计算基尼系数"""
    array = np.array(array)
    if array.size == 0:
        return 0
    array = np.sort(array)
    n = array.size
    cum_wealth = np.cumsum(array)
    gini = (2 * np.sum((np.arange(1, n+1)) * array)) / (n * np.sum(array)) - (n + 1) / n
    return gini

# 读取高道德 Excel 文件
file_path_high = r"D:\word\工作簿3高道德.xlsx"
df_high = pd.read_excel(file_path_high, sheet_name='Sheet1')

# 分组计算基尼系数和总财富
result_high = (
    df_high.groupby(['experiment', 'step'])
    .agg(
        gini=('wealth', calculate_gini),
        total_wealth=('wealth', 'sum')
    )
    .reset_index()
)

# 保存结果
output_path_high = r"D:\word\高道德_gini_total_wealth.xlsx"
result_high.to_excel(output_path_high, index=False)

print(f"高道德结果保存完成：{output_path_high}")


import pandas as pd

# 1. 文件路径和表名
file_path = r"D:\乐云云盘\thinking\merged_experiment_results.xlsx"
sheet_name = "Agent Data"

# 2. 读取数据
df = pd.read_excel(file_path, sheet_name=sheet_name)

# 3. 分组统计每一步的总财富和平均道德水平
grouped = df.groupby(["experiment", "step"]).agg(
    total_wealth=("wealth", "sum"),
    avg_moral_value=("moral_value", "mean")
).reset_index()

# 4. 计算全体 moral_value 平均值
overall_moral_mean = df["moral_value"].mean()

# 5. 添加 moral_class 列
grouped["moral_class"] = grouped["avg_moral_value"].apply(
    lambda x: "High" if x >= overall_moral_mean else "Low"
)

# 6. 创建一行用于展示总体 moral 均值
summary_row = pd.DataFrame({
    "experiment": [None],
    "step": [None],
    "total_wealth": [None],
    "avg_moral_value": [overall_moral_mean],
    "moral_class": ["Overall_Avg"]
})

# 7. 合并并写入 Excel
final_df = pd.concat([grouped, summary_row], ignore_index=True)
output_path = r"D:\乐云云盘\thinking\aggregated_moral_wealth.xlsx"
final_df.to_excel(output_path, index=False)

print(f"分析完成，结果已保存到：{output_path}")


import pandas as pd
import numpy as np

# ----------- 1. 读取数据 -----------------
file_path = r"D:\乐云云盘\thinking\merged_experiment_results.xlsx"
sheet_name = "Agent Data"
df = pd.read_excel(file_path, sheet_name=sheet_name)

# ----------- 2. 定义 Gini 系数计算函数 ---------------
def compute_gini(array):
    array = np.array(array)
    if len(array) == 0:
        return np.nan
    array = array.flatten()
    if np.amin(array) < 0:
        array -= np.amin(array)  # 所有值非负
    array += 1e-10  # 避免除以0
    array = np.sort(array)
    n = len(array)
    index = np.arange(1, n + 1)
    return ((np.sum((2 * index - n - 1) * array)) / (n * np.sum(array)))

# ----------- 3. 按 experiment 和 step 统计 -----------------
def gini_per_step(sub_df):
    return compute_gini(sub_df['wealth'])

grouped = df.groupby(["experiment", "step"]).agg(
    total_wealth=("wealth", "sum"),
    avg_moral_value=("moral_value", "mean"),
    gini_coefficient=("wealth", compute_gini)
).reset_index()

# ----------- 4. moral class 分类 -----------------
overall_moral_mean = df["moral_value"].mean()

grouped["moral_class"] = grouped["avg_moral_value"].apply(
    lambda x: "High" if x >= overall_moral_mean else "Low"
)

# ----------- 5. 总体 moral 均值作为一行 -----------------
summary_row = pd.DataFrame({
    "experiment": [None],
    "step": [None],
    "total_wealth": [None],
    "avg_moral_value": [overall_moral_mean],
    "gini_coefficient": [None],
    "moral_class": ["Overall_Avg"]
})

# ----------- 6. 保存输出 -----------------
final_df = pd.concat([grouped, summary_row], ignore_index=True)
output_path = r"D:\乐云云盘\thinking\aggregated_moral_wealth_with_gini.xlsx"
final_df.to_excel(output_path, index=False)

print(f"分析完成，结果已保存到：{output_path}")


import pandas as pd
import numpy as np

# ----------- 1. 读取数据 -----------------
file_path = r"D:\乐云云盘\thinking\merged_experiment_results.xlsx"
sheet_name = "Agent Data"
df = pd.read_excel(file_path, sheet_name=sheet_name)

# ----------- 2. 定义 Gini 系数计算函数 ---------------
def compute_gini(array):
    array = np.array(array)
    if len(array) == 0:
        return np.nan
    array = array.flatten()
    if np.amin(array) < 0:
        array -= np.amin(array)  # 所有值非负
    array += 1e-10  # 避免除以0
    array = np.sort(array)
    n = len(array)
    index = np.arange(1, n + 1)
    return ((np.sum((2 * index - n - 1) * array)) / (n * np.sum(array)))

# ----------- 3. 每组统计（按 experiment + step） -----------------
grouped = df.groupby(["experiment", "step"]).agg(
    total_wealth=("wealth", "sum"),
    gini_coefficient=("wealth", compute_gini),
    high_moral_ratio=("is_high_moral", "mean")  # True=1, False=0，求平均即为比例
).reset_index()

# ----------- 4. 计算所有组的高道德平均比例 -----------------
overall_high_moral_ratio = grouped["high_moral_ratio"].mean()

# ----------- 5. moral_class 分类 -----------------
grouped["moral_class"] = grouped["high_moral_ratio"].apply(
    lambda x: "High" if x >= overall_high_moral_ratio else "Low"
)

# ----------- 6. 总体 moral 均值作为一行 -----------------
summary_row = pd.DataFrame({
    "experiment": [None],
    "step": [None],
    "total_wealth": [None],
    "gini_coefficient": [None],
    "high_moral_ratio": [overall_high_moral_ratio],
    "moral_class": ["Overall_Avg"]
})

# ----------- 7. 合并并保存 -----------------
final_df = pd.concat([grouped, summary_row], ignore_index=True)
output_path = r"D:\乐云云盘\thinking\aggregated_moral_ratio_with_gini2.xlsx"
final_df.to_excel(output_path, index=False)

print(f"分析完成，结果已保存到：{output_path}")
